"""
CSV Merger for Code Activity and IDE Events

This script merges code activity and IDE event CSV files for each user/task combination,
sorts by timestamp, and prepares data for Excel analysis with program count and fragment diff calculations.

Usage:
python merge_csv_data.py [input_directory] [output_directory]
"""

import csv
import glob
import os
import sys
from datetime import datetime
import openpyxl
import openpyxl.styles
from openpyxl import Workbook

def merge_csv_data(input_directory, output_directory):
    """
    Main function to merge CSV files for each user/task combination
    Args:
        input_directory: Directory containing user folders with CSV files
        output_directory: Where to save merged results
    """
    # Create output directory structure
    os.makedirs(output_directory, exist_ok=True)

    # Find all user folders
    user_folders = glob.glob(os.path.join(input_directory, 'user_*'))
    print(f"Found {len(user_folders)} user folders")

    if not user_folders:
        print("No user folders found")
        return

    total_merged = 0

    for user_folder in user_folders:
        user_id = os.path.basename(user_folder)
        print(f"Processing {user_id}...")

        # Create user output directory
        user_output_dir = os.path.join(output_directory, user_id)
        os.makedirs(user_output_dir, exist_ok=True)

        # Find task folders (1, 2, 3, 4) within user folder
        task_folders = []
        for i in range(1, 5):  # Tasks 1-4
            task_folder = os.path.join(user_folder, str(i))
            if os.path.isdir(task_folder):
                task_folders.append((str(i), task_folder))

        # Process each task
        for task_id, task_folder in task_folders:
            merged_count = merge_task_files(user_id, task_id, task_folder, user_output_dir)
            total_merged += merged_count

    print(f"Successfully merged {total_merged} task combinations")


def merge_task_files(user_id, task_id, task_folder, output_dir):
    """
    Merge code activity and IDE event files for a specific task
    """
    # Find code activity files (should contain code changes/fragments)
    code_files = []
    ide_files = []

    all_csvs = glob.glob(os.path.join(task_folder, '*.csv'))

    for csv_file in all_csvs:
        filename = os.path.basename(csv_file).lower()
        if 'ide-events' in filename or 'idestate' in filename:
            ide_files.append(csv_file)
        else:
            code_files.append(csv_file)

    if not code_files or not ide_files:
        print(f"  Warning: Missing files for {user_id} task {task_id}")
        print(f"    Code files: {len(code_files)}, IDE files: {len(ide_files)}")
        return 0

    # Use the first file of each type (assuming one per task)
    code_file = code_files[0]
    ide_file = ide_files[0]

    print(f"  Merging {user_id} task {task_id}")

    # Read both CSV files
    combined_data = []

    # Read code activity data
    try:
        with open(code_file, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row in reader:
                if len(row) >= 6:  # Ensure we have enough columns
                    # Create merged row: date, fragment, source_type
                    merged_row = {
                        'date': row[0],
                        'timestamp_raw': row[0],
                        'fragment': row[5] if len(row) > 5 else '',  # Fragment column
                        'source_type': 'code_activity',
                        'raw_row': row
                    }
                    combined_data.append(merged_row)
    except Exception as e:
        print(f"  Error reading code file {code_file}: {e}")
        return 0

    # Read IDE event data
    try:
        with open(ide_file, 'r', encoding='utf-8') as f:
            reader = csv.reader(f)
            for row in reader:
                if len(row) >= 3:  # Ensure we have enough columns
                    # Create merged row: date, fragment (from Column C - index 2), source_type
                    fragment = row[2] if len(row) > 2 else ''  # Column C (index 2)

                    merged_row = {
                        'date': row[0],
                        'timestamp_raw': row[0],
                        'fragment': fragment,
                        'source_type': 'ide_event',
                        'raw_row': row
                    }
                    combined_data.append(merged_row)
    except Exception as e:
        print(f"  Error reading IDE file {ide_file}: {e}")
        return 0

    if not combined_data:
        print(f"  No data found for {user_id} task {task_id}")
        return 0

    # Create sorted version only
    sorted_filename = f"{user_id}_task{task_id}_combined.xlsx"
    sorted_path = os.path.join(output_dir, sorted_filename)
    create_excel_file(combined_data, sorted_path, sort_data=True)

    print(f"    Created: {sorted_filename}")
    return 1


def parse_timestamp(timestamp_str):
    """
    Parse timestamp string into datetime object for sorting
    More robust parsing for format: 2025-03-28T13:10:16.666-04:00
    """
    import re

    try:
        # Clean up the timestamp string
        timestamp_clean = timestamp_str.strip()

        # Use regex to extract components for your specific format
        # Pattern: YYYY-MM-DDTHH:MM:SS.sss±HH:MM
        pattern = r'(\d{4})-(\d{2})-(\d{2})T(\d{2}):(\d{2}):(\d{2})\.(\d{3})([-+])(\d{2}):(\d{2})'
        match = re.match(pattern, timestamp_clean)

        if match:
            year, month, day, hour, minute, second, microsecond, tz_sign, tz_hour, tz_min = match.groups()

            # Create datetime object (ignoring timezone for sorting consistency)
            dt = datetime(
                int(year), int(month), int(day),
                int(hour), int(minute), int(second),
                int(microsecond) * 1000  # Convert milliseconds to microseconds
            )
            return dt
        else:
            # Fallback to fromisoformat
            dt = datetime.fromisoformat(timestamp_clean)
            return dt.replace(tzinfo=None) if dt.tzinfo else dt

    except Exception as e:
        print(f"Warning: Could not parse timestamp '{timestamp_str}': {e}")
        return datetime.min  # Return minimum datetime for unparseable timestamps


def create_excel_file(data, output_path, sort_data=True):
    """
    Create Excel file with merged data and calculated columns
    """
    try:
        # Sort data if requested
        if sort_data:
            print(f"Sorting {len(data)} rows by timestamp...")

            # Add parsed timestamp for sorting
            for item in data:
                item['parsed_time'] = parse_timestamp(item['timestamp_raw'])

            # Sort by parsed timestamp
            data = sorted(data, key=lambda x: x['parsed_time'])

            # Debug: show first and last timestamps
            if data:
                print(f"First timestamp: {data[0]['timestamp_raw']}")
                print(f"Last timestamp: {data[-1]['timestamp_raw']}")

        # Create Excel workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Merged Data"

        # Headers (single row only)
        headers = ['date', 'fragment', 'program_count', 'fragment_diff']
        ws.append(headers)

        # Add data rows
        last_program_count = 0  # Track last row that had a program count

        for i, row in enumerate(data, start=2):  # Start at row 2 (after header)
            ws.cell(row=i, column=1, value=row['date'])

            # Fragment column - preserve exact formatting and handle long content
            fragment_cell = ws.cell(row=i, column=2)
            fragment_cell.value = row['fragment']
            # Enable text wrapping for long code
            fragment_cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')

            # Program count formula - only for code activity rows with actual content
            fragment = row['fragment']
            has_program_count = False
            if (row['source_type'] == 'code_activity' and
                    fragment and
                    fragment.strip()):  # Any non-empty content from code activity
                ws.cell(row=i, column=3, value=f'=LEN(B{i})')
                has_program_count = True
            else:
                ws.cell(row=i, column=3, value='')

            # Fragment diff formula - difference from previous row that had code
            if has_program_count and last_program_count > 0:
                # Find the last row with program_count
                ws.cell(row=i, column=4, value=f'=C{i}-C{last_program_count}')
            else:
                ws.cell(row=i, column=4, value='')

            # Update last_program_count if this row has one
            if has_program_count:
                last_program_count = i

        # Set column widths to handle long code
        ws.column_dimensions['A'].width = 20  # Date column
        ws.column_dimensions['B'].width = 100  # Fragment column - wide for code
        ws.column_dimensions['C'].width = 15  # Program count
        ws.column_dimensions['D'].width = 15  # Fragment diff

        # Set row heights to auto-fit content
        for row in ws.iter_rows(min_row=2):  # Skip header row
            ws.row_dimensions[row[0].row].height = None  # Auto height

        # Save workbook
        wb.save(output_path)

    except Exception as e:
        print(f"Error creating Excel file {output_path}: {e}")

def main():
    """Main function"""
    if len(sys.argv) < 3:
        print("Usage: python merge_csv_data.py [input_directory] [output_directory]")
        print("Example: python merge_csv_data.py study_data/ merged_output/")
        sys.exit(1)

    input_directory = sys.argv[1]
    output_directory = sys.argv[2]

    if not os.path.isdir(input_directory):
        print(f"Error: {input_directory} is not a valid directory")
        sys.exit(1)

    merge_csv_data(input_directory, output_directory)
    print("CSV merging complete!")

if __name__ == '__main__':
    main()