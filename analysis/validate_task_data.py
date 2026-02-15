"""
Validate Tableau Data

Reads task .xlsx files (the same ones Tableau reads), classifies actions
using the exact Tableau calculated field groupings, and prints per-student
summaries so you can verify against what Tableau displays.

Usage:
  python validate_task_data.py task1.xlsx
  python validate_task_data.py task1.xlsx task2.xlsx task3.xlsx task4.xlsx
"""

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)


# ── Tableau groupings (keep in sync with the Tableau calculated field) ──

WRITING = {
    "Edit", "EditorStartNewLine", "EditorIndentSelection",
    "EditorUnindentSelection", "EditorDeleteToWordStart",
    "Replace text", "Backspace", "Delete", "Enter", "Undo"
}

COPY_PASTE = {
    "Paste", "Paste (external)", "Copy/Paste (internal)", "Cut",
    "Accept Inline Completion", "Choose Autocomplete",
    "Choose Autocomplete Item", "EditorChooseLookupItem",
    "EditorChooseLookupItemReplace",
}

SKIP = {"Session total", "Session started", None}


def validate_file(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    print(f"\n{'='*85}")
    print(f"  {xlsx_path.name}  ({len(wb.sheetnames)} students)")
    print(f"{'='*85}")
    print(f"{'Student':<12} {'W Lines+':>8} {'W Lines-':>8} {'CP Lines+':>9} {'CP Lines-':>9} {'Ot Lines+':>9}  {'W%':>5} {'CP%':>5}")
    print("-" * 85)

    unclassified = set()

    for sname in sorted(wb.sheetnames, key=lambda s: int(s.replace("student", ""))):
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        headers = rows[0]
        ai = headers.index("action")
        try:
            lai = headers.index("lines_added")
            lri = headers.index("lines_removed")
        except ValueError:
            print(f"  {sname}: columns 'lines_added'/'lines_removed' not found — regenerate the Excel.")
            wb.close()
            return

        w_add, w_rem = 0, 0
        cp_add, cp_rem = 0, 0
        ot_add = 0
        for r in rows[1:]:
            action = r[ai]
            la = int(r[lai]) if r[lai] else 0
            lr = int(r[lri]) if r[lri] else 0
            if action == "Session total":
                continue
            elif action in WRITING:
                w_add += la
                w_rem += lr
            elif action in COPY_PASTE:
                cp_add += la
                cp_rem += lr
            else:
                ot_add += la
                if action not in SKIP:
                    unclassified.add(action)

        total_added = w_add + cp_add
        wp = w_add / total_added * 100 if total_added else 0
        cpp = cp_add / total_added * 100 if total_added else 0
        print(f"{sname:<12} {w_add:>8} {w_rem:>8} {cp_add:>9} {cp_rem:>9} {ot_add:>9}  {wp:>4.1f}% {cpp:>4.1f}%")

    wb.close()

    if unclassified:
        print(f"\nActions falling into 'Other' (not in Writing or Copy/Paste groupings):")
        for a in sorted(unclassified):
            print(f"  - {a}")


def main():
    if len(sys.argv) < 2:
        print("Usage: python validate_task_data.py task1.xlsx [task2.xlsx ...]")
        sys.exit(1)

    for arg in sys.argv[1:]:
        p = Path(arg)
        if not p.exists():
            print(f"Not found: {p}")
            continue
        validate_file(p)


if __name__ == "__main__":
    main()